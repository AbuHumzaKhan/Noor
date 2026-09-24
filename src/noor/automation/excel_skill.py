from __future__ import annotations

from pathlib import Path
from typing import Any, ClassVar


class ExcelSkill:
    """Deterministic Excel operations exposed to Noor's V1 Orchestra."""

    SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".xlsx", ".xlsm", ".xltx", ".xltm"}
    DEFAULT_MAX_ROWS: ClassVar[int] = 1000
    DEFAULT_MAX_COLUMNS: ClassVar[int] = 100
    FORMULA_OPERATIONS: ClassVar[set[str]] = {
        "sum", "average", "count", "min", "max", "if", "sumif", "countif"
    }
    TRANSFORM_OPERATIONS: ClassVar[set[str]] = {
        "remove_duplicates", "fill_blank", "rename_header"
    }

    def _validate_path(self, path: str) -> Path:
        file_path = Path(path).expanduser()
        if file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported Excel file: {file_path.suffix}")
        if not file_path.exists():
            raise FileNotFoundError(file_path)
        if not file_path.is_file():
            raise ValueError(f"Excel path is not a file: {file_path}")
        return file_path

    @staticmethod
    def _openpyxl():
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("Excel support requires openpyxl") from exc
        return openpyxl

    def inspect(self, path: str) -> dict[str, Any]:
        file_path = self._validate_path(path)
        openpyxl = self._openpyxl()
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
        try:
            sheets = [
                {"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                 "dimensions": ws.calculate_dimension()}
                for ws in workbook.worksheets
            ]
            return {"path": str(file_path), "sheets": sheets}
        finally:
            workbook.close()

    def read_sheet(
        self, path: str, sheet_name: str, max_rows: int = DEFAULT_MAX_ROWS,
        max_columns: int = DEFAULT_MAX_COLUMNS,
    ) -> dict[str, Any]:
        if max_rows < 1 or max_columns < 1:
            raise ValueError("max_rows and max_columns must be positive")
        file_path = self._validate_path(path)
        openpyxl = self._openpyxl()
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            worksheet = workbook[sheet_name]
            values = [
                list(row)
                for row in worksheet.iter_rows(
                    min_row=1, max_row=min(max_rows, worksheet.max_row), min_col=1,
                    max_col=min(max_columns, worksheet.max_column), values_only=True,
                )
            ]
            return {
                "path": str(file_path), "sheet": sheet_name, "rows_returned": len(values),
                "columns_returned": max((len(row) for row in values), default=0), "values": values,
            }
        finally:
            workbook.close()

    def write_cells(
        self, path: str, sheet_name: str, cells: dict[str, Any], output_path: str | None = None,
    ) -> dict[str, Any]:
        if not cells:
            raise ValueError("write_cells requires at least one cell")
        source = self._validate_path(path)
        destination = Path(output_path).expanduser() if output_path else source
        if destination.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported Excel output: {destination.suffix}")
        openpyxl = self._openpyxl()
        keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
        workbook = openpyxl.load_workbook(source, read_only=False, data_only=False, keep_vba=keep_vba)
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            worksheet = workbook[sheet_name]
            for coordinate, value in cells.items():
                worksheet[coordinate] = value
            workbook.save(destination)
        finally:
            workbook.close()
        return {"source": str(source), "output": str(destination), "sheet": sheet_name,
                "cells_written": len(cells), "coordinates": sorted(cells)}

    def generate_formula(
        self, operation: str, range_ref: str, *, criteria: str | None = None,
        true_value: Any = None, false_value: Any = None,
    ) -> dict[str, Any]:
        """Generate a bounded Excel formula from an allow-listed operation."""
        op = operation.strip().lower()
        if op not in self.FORMULA_OPERATIONS:
            raise ValueError(f"Unsupported formula operation: {operation}")
        if not range_ref.strip():
            raise ValueError("range_ref is required")
        ref = range_ref.strip()
        if op in {"sum", "average", "count", "min", "max"}:
            formula = f"={op.upper()}({ref})"
        elif op == "if":
            if criteria is None:
                raise ValueError("IF requires criteria")
            formula = f'=IF({criteria},{self._literal(true_value)},{self._literal(false_value)})'
        elif op in {"sumif", "countif"}:
            if criteria is None:
                raise ValueError(f"{op.upper()} requires criteria")
            formula = f'={op.upper()}({ref},{self._literal(criteria)})'
        else:
            raise ValueError(f"Unsupported formula operation: {operation}")
        return {"operation": op, "formula": formula}

    @staticmethod
    def _literal(value: Any) -> str:
        if value is None:
            return '""'
        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"
        if isinstance(value, (int, float)):
            return str(value)
        text = str(value).replace('"', '""')
        return f'"{text}"'

    def transform_sheet(
        self, path: str, sheet_name: str, operation: str, output_path: str | None = None,
        *, column: str | None = None, value: Any = None, new_name: str | None = None,
    ) -> dict[str, Any]:
        """Apply a small, deterministic transformation without executing macros."""
        op = operation.strip().lower()
        if op not in self.TRANSFORM_OPERATIONS:
            raise ValueError(f"Unsupported transformation: {operation}")
        source = self._validate_path(path)
        destination = Path(output_path).expanduser() if output_path else source
        openpyxl = self._openpyxl()
        keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
        workbook = openpyxl.load_workbook(source, read_only=False, data_only=False, keep_vba=keep_vba)
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            ws = workbook[sheet_name]
            changed = 0
            if op == "remove_duplicates":
                seen: set[tuple[Any, ...]] = set()
                delete_rows: list[int] = []
                for row in ws.iter_rows(min_row=2, values_only=False):
                    key = tuple(cell.value for cell in row)
                    if key in seen:
                        delete_rows.append(row[0].row)
                    else:
                        seen.add(key)
                for row_number in reversed(delete_rows):
                    ws.delete_rows(row_number)
                changed = len(delete_rows)
            elif op == "fill_blank":
                for row in ws.iter_rows(min_row=2):
                    cells = row if column is None else (ws[f"{column}{row[0].row}"],)
                    for cell in cells:
                        if cell.value is None:
                            cell.value = value
                            changed += 1
            elif op == "rename_header":
                if not column or new_name is None:
                    raise ValueError("rename_header requires column and new_name")
                cell = ws[f"{column}1"]
                if cell.value is None:
                    raise ValueError(f"Header not found in {column}1")
                cell.value = new_name
                changed = 1
            workbook.save(destination)
        finally:
            workbook.close()
        return {"source": str(source), "output": str(destination), "sheet": sheet_name,
                "operation": op, "changes": changed}

    def analyze_sheet(self, path: str, sheet_name: str) -> dict[str, Any]:
        """Return deterministic descriptive analytics for a worksheet."""
        file_path = self._validate_path(path)
        try:
            import pandas as pd
        except ImportError as exc:
            raise RuntimeError("Excel analysis requires pandas") from exc
        frame = pd.read_excel(file_path, sheet_name=sheet_name)
        numeric = frame.select_dtypes(include="number")
        summary: dict[str, dict[str, Any]] = {}
        for name in numeric.columns:
            series = numeric[name].dropna()
            summary[str(name)] = {
                "count": int(series.size),
                "missing": int(frame[name].isna().sum()),
                "sum": float(series.sum()) if not series.empty else 0.0,
                "mean": float(series.mean()) if not series.empty else None,
                "min": float(series.min()) if not series.empty else None,
                "max": float(series.max()) if not series.empty else None,
            }
        return {"path": str(file_path), "sheet": sheet_name, "rows": len(frame),
                "columns": len(frame.columns), "numeric_summary": summary}
