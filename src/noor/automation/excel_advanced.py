from __future__ import annotations

import re
from pathlib import Path
from typing import Any, ClassVar


class AdvancedExcelSkill:
    """Deterministic, non-macro Excel operations for V1 workflows."""

    SUPPORTED_EXTENSIONS: ClassVar[set[str]] = {".xlsx", ".xlsm", ".xltx", ".xltm"}

    def _path(self, path: str) -> Path:
        file_path = Path(path).expanduser()
        if file_path.suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported Excel file: {file_path.suffix}")
        if not file_path.is_file():
            raise FileNotFoundError(file_path)
        return file_path

    @staticmethod
    def _openpyxl():
        try:
            import openpyxl
        except ImportError as exc:
            raise RuntimeError("Excel support requires openpyxl") from exc
        return openpyxl

    def _load(self, path: str):
        source = self._path(path)
        openpyxl = self._openpyxl()
        keep_vba = source.suffix.lower() in {".xlsm", ".xltm"}
        return source, openpyxl.load_workbook(source, data_only=False, keep_vba=keep_vba)

    def workbook_summary(self, path: str) -> dict[str, Any]:
        source, workbook = self._load(path)
        try:
            return {
                "path": str(source),
                "sheet_count": len(workbook.worksheets),
                "active_sheet": workbook.active.title,
                "sheets": [
                    {"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                     "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
                     "auto_filter": ws.auto_filter.ref or None}
                    for ws in workbook.worksheets
                ],
            }
        finally:
            workbook.close()

    def search(self, path: str, query: str, sheet_name: str | None = None, limit: int = 100) -> dict[str, Any]:
        if not query.strip():
            raise ValueError("Search query cannot be empty")
        if limit < 1 or limit > 1000:
            raise ValueError("limit must be between 1 and 1000")
        source, workbook = self._load(path)
        pattern = re.compile(re.escape(query), re.IGNORECASE)
        matches: list[dict[str, Any]] = []
        try:
            worksheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets
            for ws in worksheets:
                for row in ws.iter_rows():
                    for cell in row:
                        if pattern.search(str(cell.value)):
                            matches.append({"sheet": ws.title, "cell": cell.coordinate, "value": cell.value})
                            if len(matches) >= limit:
                                return {"path": str(source), "query": query, "matches": matches, "truncated": True}
            return {"path": str(source), "query": query, "matches": matches, "truncated": False}
        finally:
            workbook.close()

    def create_sheet(self, path: str, sheet_name: str, output_path: str | None = None) -> dict[str, Any]:
        self._validate_sheet_name(sheet_name)
        source, workbook = self._load(path)
        destination = Path(output_path).expanduser() if output_path else source
        try:
            if sheet_name in workbook.sheetnames:
                raise ValueError(f"Worksheet already exists: {sheet_name}")
            workbook.create_sheet(sheet_name)
            workbook.save(destination)
            return {"path": str(destination), "sheet": sheet_name, "created": True}
        finally:
            workbook.close()

    def rename_sheet(self, path: str, sheet_name: str, new_name: str, output_path: str | None = None) -> dict[str, Any]:
        self._validate_sheet_name(new_name)
        source, workbook = self._load(path)
        destination = Path(output_path).expanduser() if output_path else source
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            if new_name in workbook.sheetnames:
                raise ValueError(f"Worksheet already exists: {new_name}")
            workbook[sheet_name].title = new_name
            workbook.save(destination)
            return {"path": str(destination), "old_name": sheet_name, "new_name": new_name}
        finally:
            workbook.close()

    def freeze(self, path: str, sheet_name: str, cell: str, output_path: str | None = None) -> dict[str, Any]:
        if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*", cell):
            raise ValueError("cell must be an Excel coordinate such as A2 or B3")
        source, workbook = self._load(path)
        destination = Path(output_path).expanduser() if output_path else source
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            workbook[sheet_name].freeze_panes = cell
            workbook.save(destination)
            return {"path": str(destination), "sheet": sheet_name, "freeze_panes": cell}
        finally:
            workbook.close()

    def set_filter(self, path: str, sheet_name: str, cell_range: str, output_path: str | None = None) -> dict[str, Any]:
        if not re.fullmatch(r"[A-Za-z]{1,3}[1-9][0-9]*:[A-Za-z]{1,3}[1-9][0-9]*", cell_range):
            raise ValueError("cell_range must look like A1:F100")
        source, workbook = self._load(path)
        destination = Path(output_path).expanduser() if output_path else source
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            workbook[sheet_name].auto_filter.ref = cell_range
            workbook.save(destination)
            return {"path": str(destination), "sheet": sheet_name, "auto_filter": cell_range}
        finally:
            workbook.close()

    def sort_by_column(self, path: str, sheet_name: str, column: str, output_path: str | None = None, descending: bool = False) -> dict[str, Any]:
        if not re.fullmatch(r"[A-Za-z]{1,3}", column):
            raise ValueError("column must be an Excel column such as A or AA")
        source, workbook = self._load(path)
        destination = Path(output_path).expanduser() if output_path else source
        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"Worksheet not found: {sheet_name}")
            ws = workbook[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) <= 2:
                workbook.save(destination)
                return {"path": str(destination), "sheet": sheet_name, "rows_sorted": max(len(rows) - 1, 0)}
            index = self._column_index(column)
            header, data = rows[0], rows[1:]
            data.sort(key=lambda row: (row[index] is None, str(row[index]).casefold()), reverse=descending)
            for row_number, values in enumerate([header, *data], start=1):
                for col_number, value in enumerate(values, start=1):
                    ws.cell(row=row_number, column=col_number, value=value)
            workbook.save(destination)
            return {"path": str(destination), "sheet": sheet_name, "rows_sorted": len(data), "column": column, "descending": descending}
        finally:
            workbook.close()

    @staticmethod
    def validate_formula(formula: str) -> dict[str, Any]:
        value = formula.strip()
        if not value.startswith("="):
            return {"valid": False, "formula": formula, "reason": "Formula must start with '='"}
        if value.count("(") != value.count(")"):
            return {"valid": False, "formula": formula, "reason": "Unbalanced parentheses"}
        if any(token in value.upper() for token in ("WEBSERVICE(", "HYPERLINK(", "CALL(", "EXEC(")):
            return {"valid": False, "formula": formula, "reason": "Formula contains a blocked external/action function"}
        return {"valid": True, "formula": value, "reason": "Basic syntax checks passed"}

    @staticmethod
    def explain_formula(formula: str) -> dict[str, Any]:
        value = formula.strip()
        validation = AdvancedExcelSkill.validate_formula(value)
        if not validation["valid"]:
            return validation
        function = re.match(r"=([A-Z][A-Z0-9_.]*)\s*\(", value.upper())
        name = function.group(1) if function else "expression"
        descriptions = {
            "SUM": "Adds numeric values or ranges.", "AVERAGE": "Calculates the arithmetic mean.",
            "COUNT": "Counts numeric cells.", "IF": "Returns one value when a condition is true and another when false.",
            "SUMIF": "Adds cells that satisfy a criterion.", "COUNTIF": "Counts cells that satisfy a criterion.",
            "XLOOKUP": "Looks up a value and returns a corresponding result.", "VLOOKUP": "Looks up a value in the first column of a table.",
        }
        return {"valid": True, "formula": value, "function": name, "explanation": descriptions.get(name, "Excel expression; inspect its arguments for the exact calculation.")}

    @staticmethod
    def _validate_sheet_name(name: str) -> None:
        if not name.strip() or len(name) > 31 or re.search(r"[\\/:*?\[\]]", name):
            raise ValueError("Invalid Excel worksheet name")

    @staticmethod
    def _column_index(column: str) -> int:
        result = 0
        for char in column.upper():
            result = result * 26 + ord(char) - 64
        return result - 1
