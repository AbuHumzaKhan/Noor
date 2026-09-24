from pathlib import Path

import openpyxl

from noor.automation.excel_skill import ExcelSkill
from noor.automation.registry import default_registry


def _workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sales"
    sheet.append(["Product", "Sales", "Region"])
    sheet.append(["A", 100, "North"])
    sheet.append(["B", 200, "South"])
    sheet.append(["B", 200, "South"])
    workbook.save(path)


def test_formula_generation() -> None:
    result = ExcelSkill().generate_formula("sum", "B2:B10")
    assert result["formula"] == "=SUM(B2:B10)"


def test_formula_rejects_unknown_operation() -> None:
    try:
        ExcelSkill().generate_formula("execute", "B2:B10")
    except ValueError as exc:
        assert "Unsupported formula operation" in str(exc)
    else:
        raise AssertionError("Unknown formula operation should fail closed")


def test_transform_removes_duplicate_rows(tmp_path: Path) -> None:
    source = tmp_path / "sales.xlsx"
    output = tmp_path / "clean.xlsx"
    _workbook(source)
    result = ExcelSkill().transform_sheet(str(source), "Sales", "remove_duplicates", str(output))
    assert result["changes"] == 1
    workbook = openpyxl.load_workbook(output, read_only=True)
    assert workbook["Sales"].max_row == 3
    workbook.close()


def test_analyze_sheet(tmp_path: Path) -> None:
    source = tmp_path / "sales.xlsx"
    _workbook(source)
    result = ExcelSkill().analyze_sheet(str(source), "Sales")
    assert result["rows"] == 3
    assert result["numeric_summary"]["Sales"]["sum"] == 500.0


def test_registry_exposes_v1_excel_tasks() -> None:
    names = {spec.name for spec in default_registry().list()}
    assert {"excel.formula.generate", "excel.transform", "excel.analyze"} <= names
